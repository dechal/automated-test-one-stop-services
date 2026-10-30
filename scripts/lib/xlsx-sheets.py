import json
import sys
from datetime import date, datetime

from openpyxl import load_workbook


def cell_text(value: object) -> str:
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    return str(value)


def main() -> int:
    wb = load_workbook(sys.argv[1], data_only=True)
    sheets = []
    for ws in wb.worksheets:
        width = ws.max_column or 0
        rows = [
            [cell_text(ws.cell(row=r, column=c).value) for c in range(1, width + 1)]
            for r in range(1, (ws.max_row or 0) + 1)
        ]
        sheets.append({'name': ws.title, 'rows': rows})
    print(json.dumps(sheets, ensure_ascii=False))
    return 0


if __name__ == '__main__':
    sys.exit(main())
